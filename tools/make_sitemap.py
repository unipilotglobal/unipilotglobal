#!/usr/bin/env python3
"""UNIPILOT GLOBAL: builds sitemap.xml (the list of pages that Google should know about) from your content.xlsx.

How to use it by hand (only if you do not use the automatic GitHub step):
    pip install openpyxl
    python tools/make_sitemap.py content.xlsx sitemap.xml
Then upload the new sitemap.xml with GitHub Desktop like any other file.

It lists: the main pages, every opportunity that is not marked Hidden, and every mentor in the Experts sheet that is not Draft or Hidden.
"""
import re, sys, os, json, datetime
from urllib.parse import quote
from xml.sax.saxutils import escape

try:
    import openpyxl
except ImportError:
    sys.exit('Please run:  pip install openpyxl')

SITE = 'https://unipilotglobal.com'
PAGES = [('', 1.0), ('scholarships.html', 0.9), ('opportunities.html', 0.9), ('pilots.html', 0.8), ('services.html', 0.7), ('materials.html', 0.7),
         ('seminars.html', 0.6), ('news.html', 0.6), ('reviews.html', 0.5), ('contact.html', 0.5), ('legal.html', 0.3)]


def slug(s):                       # the same rule the website uses to build an address from a name
    return re.sub(r'^-|-$', '', re.sub(r'[^a-z0-9]+', '-', str(s).lower()))[:70]


def rows(wb, name):
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    head = next(it, None)
    if not head:
        return []
    head = [str(h).strip() if h is not None else '' for h in head]
    return [dict(zip(head, r)) for r in it if r and any(c not in (None, '') for c in r)]


def static_ids(out):               # opportunities that have a ready-made page in /o/ (written by tools/prerender.js)
    f = os.path.join(os.path.dirname(os.path.abspath(out)), 'data', 'static-ids.js')
    try:
        return set(json.loads(re.search(r'UP_STATIC\s*=\s*(\[.*?\])', open(f, encoding='utf-8').read(), re.S).group(1)))
    except Exception:
        return set()


def main(src='content.xlsx', out='sitemap.xml'):
    today = datetime.date.today().isoformat()
    ready = static_ids(out)
    urls = [(SITE + '/' + p, today, pr) for p, pr in PAGES]
    seen = set()
    if src:
        wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
        for r in rows(wb, 'Opportunities'):
            name = str(r.get('Name') or '').strip()
            if not name or str(r.get('Status') or '').strip().lower() == 'hidden':
                continue
            oid = str(r.get('ID') or '').strip() or slug(name)
            if not oid or oid in seen:
                continue
            seen.add(oid)
            v = str(r.get('Verified on') or '')[:10]
            page = ('/o/' + quote(oid, safe='-_.~') + '.html') if oid in ready else ('/opportunity.html?id=' + quote(oid, safe='-_.~'))
            urls.append((SITE + page, v if re.match(r'^\d{4}-\d{2}-\d{2}$', v) else today, 0.6))
        for r in rows(wb, 'Experts'):
            name = str(r.get('Name') or '').strip()
            if not name or re.match(r'^(draft|hidden)$', str(r.get('Status') or '').strip(), re.I):
                continue
            urls.append((SITE + '/pilots.html?pilot=' + quote(slug(name), safe='-_.~'), today, 0.5))
    lines = ['<?xml version="1.0" encoding="UTF-8"?>', '<urlset xmlns="http://www.sitemaps.org/schemas/sitemap/0.9">']
    for loc, mod, pr in urls:
        lines.append('<url><loc>%s</loc><lastmod>%s</lastmod><priority>%.1f</priority></url>' % (escape(loc), mod, pr))
    lines.append('</urlset>')
    open(out, 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('Wrote %s with %d addresses (%d opportunities).' % (out, len(urls), len(seen)))


if __name__ == '__main__':
    a = sys.argv[1:]
    if a and a[0] == '--pages-only':
        main(None, a[1] if len(a) > 1 else 'sitemap.xml')
    else:
        main(a[0] if a else 'content.xlsx', a[1] if len(a) > 1 else 'sitemap.xml')
